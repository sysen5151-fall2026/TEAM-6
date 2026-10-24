"""Builds EcoCommute_Team6_Ch6-Ch7_Architecture_Definition_Function.docx"""
import json
import os
import re
import sys
import docx as _docx
import figs3
import mdata as M
from docx_helpers import *

HERE = os.path.dirname(os.path.abspath(__file__))
DIA = os.path.join(HERE, "..", "diagrams")
OLD_DIA = os.path.join(HERE, "..", "..", "ch4-ch5", "diagrams")
OLD_DOC = os.path.join(HERE, "..", "..", "ch4-ch5", "EcoCommute_Team6_Ch4-Ch5_System_Requirements_Definition_Final.docx")
REQS = json.load(open(os.environ.get("REQS_JSON", os.path.join(HERE, "reqs.json")), encoding="utf8"))
OUT = os.path.join(HERE, "..", "EcoCommute_Team6_Ch6-Ch7_Architecture_Definition_Function.docx")

# ------------------------------------------------------------------ numbering registry
HAS_F2 = os.path.exists(os.path.join(DIA, "innoslate_hierarchy_f2.png"))
FIG_ORDER = ["universe", "meta", "a0", "first", "hier"] + (["part1"] if HAS_F2 else []) + ["n2",
             "action", "metatrace", "innoseq", "seq", "trace", "matrix"]
TAB_ORDER = ["items", "funcs", "ios", "nfr", "metaicom", "a0icom", "purpose", "icom1", "n2flows",
             "actmap", "seqmap", "findings", "changes", "freq", "cover", "tracecrit", "idefcrit"]
FIG = {k: i + 1 for i, k in enumerate(FIG_ORDER)}
TAB = {k: i + 1 for i, k in enumerate(TAB_ORDER)}
_fig_seen, _tab_seen = [], []


def F(k):
    return f"Figure {FIG[k]}"


def T(k):
    return f"Table {TAB[k]}"


d = new_document()


def add_figure(key, path, width, text):
    _fig_seen.append(key)
    assert FIG_ORDER[len(_fig_seen) - 1] == key, f"figure order: {key}"
    picture(d, path, width)
    caption(d, f"Figure {FIG[key]}. {text}")


def add_table(key, text, header, rows, widths, **kw):
    _tab_seen.append(key)
    assert TAB_ORDER[len(_tab_seen) - 1] == key, f"table order: {key}"
    cap = caption(d, f"Table {TAB[key]}. {text}")
    cap.paragraph_format.keep_with_next = True
    return table(d, header, rows, widths, **kw)


def strip(s):
    return re.sub(r"<[^>]+>", "", s or "").strip()


def parse_req(num):
    r = next(v for v in REQS.values() if v["num"] == num)
    rat = r["Rationale"]
    out = {"stmt": strip(r["desc"]), "name": r["name"]}
    for key, lab in (("rationale", "Rationale:"), ("mop", "MOP:"), ("mos", "MOS:"), ("ver", "Verification criteria:"), ("srd", "Stakeholder requirements:")):
        m = re.search(re.escape(lab) + r"\s*(.*?)(?:\n|$)", rat)
        out[key] = m.group(1).strip() if m else ""
    return out


old = _docx.Document(OLD_DOC)


def old_table(i):
    t = old.tables[i]
    return [[c.text.replace("\n", " ") for c in r.cells] for r in t.rows]


# ================================================================== title + contribution
for txt, sz, b in [("SYSEN 5151 Foundations of Systems Engineering", 12, False), ("EcoCommute (Team 6)", 22, True),
                   ("Architecture Definition: Function", 18, True),
                   ("Tracing Use Cases (Chapter 6) and First-Level Function Architecture (Chapter 7)", 13, False),
                   ("Report section combining the Chapter 6 and Chapter 7 assignments", 10.5, False),
                   ("Innoslate model: EcoCommute 10.8 (10.7 export of 24 October 2026 plus the two trace edits in Section 2.1.4.1.3)   |   October 2026", 9.5, False)]:
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(txt)
    r.font.size = Pt(sz)
    r.bold = b
    if b:
        r.font.color.rgb = NAVY

heading(d, "Team Contribution", 1)
para(d, "The table lists who contributed which part of this assignment. Roles come from the team charter in the repository README.")
table(d, ["Team member", "Role", "Contribution to Chapters 6 and 7"], [
    ["Luyu Chen", "Project Manager and Integration Lead", "Integrated the report and the Innoslate model (EcoCommute 10.8); Section 3 requirements traceability matrix; final check that every figure and table is explained once and only once."],
    ["Mingjie Fang", "Quality, Risk and Documentation Lead", "Tracing Assessment Criteria and IDEF Assessment Criteria checks (Appendix A); findings and observations of the trace (Section 2.1.4.1.3); report editing."],
    ["Karen Li", "Product and Stakeholder Lead", "OpsCon and action diagram numbering (Section 2.1.4.1.1); UC.1 sequence; rationale and traceability of the functional requirements."],
    ["Yinping Yin", "Sustainability Data and Analysis Lead", "Functional requirement table with MOP and verification criteria (Section 3.1.1); non-functional elements and their influence on the functions."],
    ["Yiyang Zou", "Technical Architecture Lead", "Metafunction, IDEF0-A0 and first-level IDEF0; hierarchy diagrams; N2 diagram; sequence diagram and traced IDEF0 figures."],
], [1.2, 1.7, 3.8], size=8.5, first_bold=True)

# TOC
p = para(d, "**Table of Contents**", size=13)
toc = d.add_paragraph()
add_field(toc, 'TOC \\o "1-5" \\h \\z \\u', "Right-click and choose Update Field to refresh the table of contents.")
d.add_page_break()

# ================================================================== 1 Introduction
heading(d, "1 Introduction", 1)
heading(d, "1.1 Purpose of Architecture Definition", 2)
para(d, "Architecture Definition turns the system requirements into a model of what the system does and how its parts fit together, so that the design can be allocated and verified later. For the function view it answers four questions: which functions does EcoCommute perform, how are they organised, how do they exchange information, and does that structure really support the operational concept (OpsCon)? Chapter 6 checks the second and fourth questions by tracing the use case through the IDEF0 diagrams; Chapter 7 builds the first-level function architecture that the trace runs over. The result is a function model that the team can check against the OpsCon before it is mapped to form (physical) elements in the next chapter.")
heading(d, "1.2 Description of Architecture Definition", 2)
para(d, "The system of interest (SoI) is C.0 EcoCommute, a responsive web application that estimates commute CO2e and suggests feasible greener options. Its external assets are X.1 Commuter, X.2 Pilot Administrator, X.3 Open-Meteo, X.4 Emission Factor Publisher and X.5 Hosting Provider. The scenario traced in this report is the nominal use case UC.1, Review Commute Impact and Feasible Alternative, which assumes an existing account and consent, valid inputs, suitable stored emission factors and an available forecast.")
para(d, f"This section builds on the Chapter 4 and 5 report. The Metafunction, IDEF0-A0 and first-level IDEF0 of that report are carried forward and corrected where the trace required it. New in this section are the function hierarchy diagram, the N2 diagram, the trace of UC.1 through the Metafunction and the first-level IDEF0, the functional requirements table and the function-to-requirement traceability matrix. {T('items')} lists the assignment items and where each is covered.")
add_table("items", "Assignment items and where they are covered in this report.",
          ["Item", "Activity", "Source", "Section"], [
              ["Function architecture", "Create function hierarchy diagrams", "Innoslate model 10.8", "2.1.2.2"],
              ["Function architecture", "Create IDEF0 diagrams", "Innoslate model 10.8", "2.1.2.1"],
              ["Function interface", "Create N2 diagram", "Innoslate model 10.8", "2.1.3.1"],
              ["Architecture tracing", "Describe the tracing process that verifies the architecture structure", "Chapter 6", "2.1.4.1"],
              ["Functional requirements", "State the system functional requirements with rationale, MOP and verification criteria", "Innoslate model 10.8", "3.1.1"],
              ["Requirements traceability", "Trace requirements to functions", "Innoslate model 10.8", "3.1.2"],
              ["Team contribution", "State which student did which part", "Team", "Front matter"],
          ], [1.6, 3.2, 1.5, 0.8], size=8.5)
para(d, "**Assumptions and conventions.**", after=2)
bullet(d, "The model baseline is the Innoslate export EcoCommute_10.7_C4-C5_system_requirements of 24 October 2026. The trace exposed two inconsistencies; the edits that remove them make up model version 10.8 (Section 2.1.4.1.3).")
bullet(d, "Function names are verb phrases and input/output names are noun phrases. Numbering is F.0 for the Metafunction, F.1 to F.6 for the top-level functions and F.2.0 to F.2.7 for the first-level functions. Input/outputs are numbered IO.1 to IO.22.")
bullet(d, "First-level functions have no mechanisms yet; these are allocated in the form (physical) architecture chapter, as the lab instructs.")
bullet(d, "All model diagrams are Innoslate screenshots of model 10.8. The coloured lines and number badges of the two traced figures were added to those screenshots outside Innoslate, as the lab prescribes for tracing. The tracing sequence diagram (Section 2.1.4.1.2) is the only figure drawn outside Innoslate, from the model data, because the Innoslate sequence diagram is incomplete and cannot show the internal first-level interactions; the Innoslate version is included next to it.")
bullet(d, "Values taken from the stakeholder requirements are used unchanged; the 3-second display time and the minimum group size of 5 remain TBD pending sponsor approval, as in the previous report.")

# ================================================================== 2 Architecture Definition
heading(d, "2 Architecture Definition", 1)
heading(d, "2.1 Function Architecture", 2)
heading(d, "2.1.1 Function Architecture Model", 3)
heading(d, "2.1.1.1 Functional Elements and Non-Functional Elements", 4)
para(d, f"**Functional elements.** The function architecture has 15 functions on three levels and 22 input/outputs. {T('funcs')} lists the functions with the asset that performs them. The six top-level functions are performed by one asset each, as the Metafunction requires. The eight first-level functions decompose F.2 and are allocated to mechanisms later.")
rows = [["F.0", "Metafunction", "Universe", "0", "Holds the top-level function of the SoI and of every external asset."]]
descr = {"F.1": "Records commutes, states consent and settings, uses the advice.", "F.2": "Turns entries, forecasts and reviewed factors into an estimate, baseline comparison, trend, feasible alternative and approved summaries.",
         "F.3": "Sets reporting rules, requests and receives approved summaries.", "F.4": "Returns a forecast for an approved coarse location and trip time.",
         "F.5": "Releases reviewed emission factors with source, units, geography and version.", "F.6": "Runs the service and reports hosting status."}
for f, n in M.TOP.items():
    rows.append([f, n, M.MECH[f], "1", descr[f]])
l1d = {"F.2.0": "Deploy releases, monitor availability, run retirement.", "F.2.1": "Present the entry form, validate trip details, enforce consent, route deletions.",
       "F.2.2": "Hold reviewed factors, select the factor, obtain the forecast.", "F.2.3": "Estimate trip CO2e and compare it with the always-drive baseline.",
       "F.2.4": "Store trips, update trends, restrict access, delete on request.", "F.2.5": "Select one feasible alternative with a reason.",
       "F.2.6": "Assemble and display the result and explanation.", "F.2.7": "Aggregate usage and CO2e into approved summaries."}
for f, n in M.L1:
    rows.append([f, n, "not yet allocated", "2", l1d[f]])
add_table("funcs", "Functional elements: the 15 functions of the architecture.", ["ID", "Function", "Performed by (mechanism)", "Level", "Role"], rows, [0.55, 1.8, 1.4, 0.5, 2.85], size=8.2, first_bold=True)

io_rows = []
for io, nm in M.IO.items():
    n = int(io.split(".")[1])
    srcs, dests = M.L1FLOW[io]
    if n <= 11:
        s, dd, role = M.META[io]
        kind = {"I": "Input", "C": "Control", "O": "Output"}[role]
        if role == "O":
            kind = "Output"
        origin = f"Carried down: {s} to {dd}"
    else:
        origin = "Internal to F.2"
        rl = {}
        for dd, r_ in dests:
            rl.setdefault(r_, []).append(dd)
        if len(rl) == 1:
            kind = "Control" if "C" in rl else "Input"
        else:
            kind = "Input; control at " + ", ".join(rl["C"])
    srcs_s = ", ".join(x for x in srcs if x != "ENV") or "outside F.2"
    dst_s = ", ".join(x for x, _ in dests if x != "ENV") or "outside F.2"
    io_rows.append([io, nm, kind, origin, f"{srcs_s} to {dst_s}"])
add_table("ios", "Functional elements: the 22 input/outputs. IO.1 to IO.11 are carried down from the Metafunction; IO.12 to IO.22 are internal to F.2.",
          ["ID", "Input/output", "Kind at destination", "Origin", "First-level source to destination"], io_rows, [0.5, 1.9, 1.2, 1.7, 1.8], size=8, first_bold=True)
para(d, "**Non-functional elements.** The system-wide requirements (3.4.x) are not functions, but they decide how some functions must be built, so they were tagged to the functions they influence most. The statements and thresholds are those in the model.")
nfr_map = {
    "3.4.1.1": ("F.2 as a whole", "Keeps the decomposition to eight functions and rules out extra subsystems or services before the form architecture."),
    "3.4.1.2": ("F.2.4, F.2.7", "Trip records and pilot summaries must be collected from the first day so that results can be reported on time."),
    "3.4.2.1": ("F.2.2, F.2.3, F.2.6", "The forecast call (Weather Request, IO.4) is the only external call in the UC.1 path, so its latency dominates the display time."),
    "3.4.2.2": ("F.2.1", "Entry validation must stay light: few fields and one pass, so a commute can be entered within 30 seconds."),
    "3.4.3.1": ("F.2.0, F.2.2", "F.2.0 monitors availability (IO.22); F.2.2 must degrade gracefully when the forecast service is unavailable."),
    "3.4.3.2": ("F.2.4, F.2.0", "Record storage and the hosting configuration (IO.10) are sized for at least 500 accounts."),
    "3.4.4.1": ("F.2.1, F.2.6", "The two user-facing functions carry the accessibility requirement."),
    "3.4.5.1": ("F.2.4, F.2.7", "Only F.2.4 holds identifiable trips; F.2.7 receives a de-identified extract (IO.19)."),
    "3.4.5.2": ("F.2.1, F.2.2, F.2.4", "Only a coarse area is accepted, stored and sent in the Weather Request."),
    "3.4.5.3": ("F.2.7", "The suppression rule is a control (IO.9) applied before any summary leaves F.2.7."),
    "3.4.6.1": ("F.2.2, F.2.3", "The factor record (IO.14) and the estimate must be reproducible, so factor version and distance convention are carried with the result."),
}
nrows = []
for num, (fn, why) in nfr_map.items():
    r = parse_req(num)
    thr = r["mos"] or r["rationale"]
    nrows.append([num, r["stmt"].replace("The EcoCommute system shall ", "Shall ").rstrip("."), fn, why])
add_table("nfr", "Non-functional elements: system-wide requirements and the functions they influence.",
          ["ID", "Requirement", "Functions most affected", "How it shapes the functions"], nrows, [0.6, 2.3, 1.2, 3.0], size=8, first_bold=True)

heading(d, "2.1.2 Function Structure", 3)
heading(d, "2.1.2.1 IDEF0 Diagram", 4)
para(d, f"**Universe and mechanisms.** The Universe hierarchy ({F('universe')}) shows the SoI C.0 and the five external assets as children of the Universe, with the internal components C.1 to C.4 under C.0. Each box on the Metafunction has exactly one mechanism from this hierarchy.")
add_figure("universe", os.path.join(DIA, "trim_universe_hierarchy.png"), 6.9, "Universe hierarchy: SoI C.0 EcoCommute, external assets X.1 to X.5 and components C.1 to C.4.")
para(d, f"**Metafunction.** The Metafunction ({F('meta')}) holds one function for the SoI (F.2) and one for each external asset (F.1, F.3 to F.6): six functions, inside the five-to-nine guideline. Input/outputs are deliberately broad bundles, so that the diagram stays stable when lower-level decisions change. {T('metaicom')} gives the ICOM of each function. Every input/output has a source and a destination, so there are no tunnels. Controls are the inputs that govern how F.2 behaves rather than what it transforms: consent and settings (IO.2), reviewed emission factors (IO.6) and pilot program rules (IO.9). The external functions have implicit controls outside our design scope.")
add_figure("meta", os.path.join(DIA, "trim_metafunction_idef0.png"), 6.4, "Metafunction IDEF0 (top level) of EcoCommute from Innoslate.")
mt = old_table(2)
add_table("metaicom", "Top-level functions, mechanisms and ICOMs (Metafunction).", mt[0], mt[1:], [1.35, 1.0, 1.7, 0.95, 2.1], size=7.8, first_bold=True)
para(d, f"**IDEF0-A0.** The IDEF0-A0 diagram ({F('a0')}) is the black-box view of F.2 as seen from the Metafunction: four inputs, three controls, four outputs and the mechanism C.0. {T('a0icom')} describes every ICOM. Reading it from the Commuter's side: entries and requests (IO.1) and the controlling consent and settings (IO.2) go in, and the commute result and advice (IO.3) comes out. The forecast loop is the weather request (IO.4) out and the forecast data (IO.5) back. The pilot loop is the summary request (IO.7) and program rules (IO.9) in, and the approved aggregate summary (IO.8) out. Operations are the hosting status (IO.11) in and the release configuration (IO.10) out; the reviewed emission factors (IO.6) govern the estimate.")
add_figure("a0", os.path.join(DIA, "trim_a0_icom.png"), 5.6, "IDEF0-A0 (ICOM) diagram of F.2 Provide Commute Impact Decision Support with mechanism C.0 EcoCommute.")
a0 = old_table(3)
for row in a0[1:]:
    if row[1].startswith("IO.3 "):
        row[3] = "Entry form and validation feedback presented to the Commuter before the result (from F.2.1), then estimated CO2e, trend, baseline difference, method disclosure, one feasible alternative and plain-language explanation (from F.2.6). Updated by the trace (edit 1)."
add_table("a0icom", "ICOM description for F.2 (IDEF0-A0).", a0[0], a0[1:], [0.75, 2.0, 0.8, 3.55], size=7.8)

# first-level IDEF0 on a landscape page
para(d, f"**First-level IDEF0.** {F('first')} decomposes F.2 into eight functions. The input/outputs of the Metafunction are carried down and allocated to the functions that use or produce them; the internal flows IO.12 to IO.22 connect the functions. Inputs enter from the left, controls from the top and outputs leave on the right. F.2.1 to F.2.7 follow the data chain down the diagram and F.2.0 is placed last. Commute Result and Advice (IO.3) leaves from both F.2.1 and F.2.6.", size=10, after=2)
add_figure("first", os.path.join(DIA, "innoslate_first_level_idef0.png"), 7.0, "First-level IDEF0 of F.2 from Innoslate (model 10.8). Eight functions, 11 carried-down and 11 internal input/outputs, no tunnels. Controls enter at the top of a box.")
para(d, f"**First-level functions.** {T('purpose')} gives, for each first-level function, its goal, its key actions and the use case or OpsCon action it realizes. {T('icom1')} lists its ICOM. Reading the two together, the functions form five groups. **Intake** (F.2.1) accepts and validates entries. **Reference and analysis** (F.2.2, F.2.3, F.2.5) obtain the factor and forecast and compute the estimate and the alternative. **Records** (F.2.4) keep trips and trends. **Output** (F.2.6) assembles what the Commuter sees. **Pilot operations** (F.2.0, F.2.7) deploy the service and report on it. How the groups hand information to each other is analysed with the N2 diagram in Section 2.1.3.1 and checked with the trace in Section 2.1.4.1.")
purpose = [
    ["F.2.0 Operate Service", "Keep the pilot service running and retire it cleanly.", "Deploy releases to the hosting service (IO.10); monitor availability (IO.11) and record it (IO.22); issue the retention directive at retirement (IO.21) under the program rules (IO.9).", "OP.2, OP.4, OP.5"],
    ["F.2.1 Collect Commute Entries", "Turn a Commuter's submission into a valid, consented entry.", "Present the entry form and feedback (IO.3); accept entries and requests (IO.1); check consent and settings (IO.2); validate trip details and issue the validated entry (IO.12); route deletion requests (IO.13).", "UC.1.2, UC.1.4, UC.1.5, OP.1"],
    ["F.2.2 Manage Reference Data", "Supply the factor and forecast that govern the estimate and the advice.", "Hold the reviewed factors (IO.6); select the applicable factor record (IO.14); send the weather request with coarse area and trip time (IO.4); accept the forecast (IO.5); pass the trip forecast on (IO.15).", "UC.1.5, UC.1.6, UC.1.7"],
    ["F.2.3 Estimate Commute Emissions", "Compute trip CO2e and its difference from the always-drive baseline, reproducibly.", "Compute CO2e from distance and the factor record; compute the baseline difference; issue the estimate (IO.16) and the estimation method disclosure (IO.17).", "UC.1.8, UC.1.9"],
    ["F.2.4 Manage Commute Records", "Keep trip records and trends current and protected.", "Store the trip when consent allows (IO.12, IO.16, IO.2); update weekly and monthly trends (IO.18); restrict access; delete on request (IO.13) and at retirement (IO.21); issue a de-identified extract (IO.19).", "UC.1.8, UC.1.9, OP.1, OP.4"],
    ["F.2.5 Evaluate Feasible Alternatives", "Recommend one feasible lower-emission mode with a reason.", "Exclude modes marked infeasible (IO.2); apply distance and weather rules to the entry and trip forecast (IO.12, IO.15); select one alternative and its reason (IO.20).", "UC.1.10"],
    ["F.2.6 Present Commute Results", "Show a readable, honest result.", "Assemble estimate (IO.16), baseline difference, trend (IO.18), alternative (IO.20) and method disclosure (IO.17) with a plain-language explanation and display them (IO.3).", "UC.1.11, UC.1.12"],
    ["F.2.7 Summarize Pilot Results", "Give the sponsor pilot evidence without exposing individuals.", "Accept the summary request (IO.7); combine the trip extract (IO.19) and service status record (IO.22); apply the rules including small-group suppression (IO.9); issue the approved aggregate summary (IO.8).", "UC.2, OP.3"],
]
add_table("purpose", "First-level functions: goal, key actions and the use case or OpsCon action realized.", ["Function", "Goal", "Key actions", "Realizes"], purpose, [1.2, 1.5, 3.5, 0.9], size=8, first_bold=True)


def names(ids):
    return "; ".join(f"{i} {M.IO[i]}" for i in ids) or "none"


icom = []
for f, n in M.L1:
    ins, cts, outs = [], [], []
    for io, (srcs, dests) in M.L1FLOW.items():
        for dd, role in dests:
            if dd == f:
                (ins if role == "I" else cts).append(io)
        if f in srcs:
            outs.append(io)
    key = lambda x: int(x.split(".")[1])
    icom.append([f"{f} {n}", names(sorted(ins, key=key)), names(sorted(cts, key=key)), names(sorted(outs, key=key))])
add_table("icom1", "First-level functions and their ICOMs (model 10.8). Mechanisms are allocated in the next chapter.", ["Function", "Inputs", "Controls", "Outputs"], icom, [1.4, 2.2, 1.6, 1.9], size=7.8, first_bold=True)
para(d, "The first-level IDEF0 passes the IDEF checks: every function has a verb-oriented name and a numbered ID; every input/output is a noun phrase with a source and a destination; no first-level function has a mechanism yet; and each function has the inputs, controls and outputs it needs for its goal (Appendix A).")

heading(d, "2.1.2.2 Hierarchy Diagram", 4)
para(d, f"The function hierarchy ({F('hier')}) has three levels. Level 0 is the Metafunction F.0. Level 1 has the six top-level functions F.1 to F.6, one for the SoI and one for each external asset. Level 2 has the eight first-level functions F.2.0 to F.2.7 under F.2. Only F.2 is decomposed, because it is the function of the system we are designing; the external functions are performed by assets outside our design scope and stay as leaves. The tree therefore holds 1 + 6 + 8 = 15 functions, and F.2's eight children sit inside the five-to-nine guideline.")
add_figure("hier", os.path.join(DIA, "innoslate_hierarchy_overall.png"), 7.0, "Function hierarchy diagram from Innoslate, three levels.")
para(d, "The tree is small enough to read as one figure. The key actions of each of the eight first-level functions, which a branch diagram would show, are given in " + T('purpose') + ". The numbering follows the order in which the functions meet a commute, except that F.2.0 is numbered first so that operations come before the data chain. The five groups of first-level functions (intake, reference and analysis, records, output, pilot operations) are described in Section 2.1.2.1.")

if HAS_F2:
    para(d, f"{F('part1')} partitions the tree: it shows F.2 with its eight children, the part of the hierarchy that holds the lower-level elements, at a larger size.")
    add_figure("part1", os.path.join(DIA, "innoslate_hierarchy_f2.png"), 7.0, "Function hierarchy rooted at F.2 (Innoslate): the eight first-level functions.")
heading(d, "2.1.3 Function Interface", 3)
heading(d, "2.1.3.1 N2 Diagram", 4)
para(d, f"{F('n2')} is the N2 diagram of the first-level functions from Innoslate. The functions sit on the diagonal. A cell shows the input/outputs that flow from the function in its row (source) to the function in its column (destination). The top row above the diagram lists what enters each function from outside F.2, and the right-hand column lists what leaves to the environment; together they match the carried-down flows of the first-level IDEF0 ({F('first')}) and {T('icom1')}. The cells inside the matrix hold the 11 internal input/outputs and match the internal flows of the IDEF0 one to one. The cell text of the Innoslate export is small, so {T('n2flows')} writes out the content of every row.", size=10, after=2)
add_figure("n2", os.path.join(DIA, "innoslate_n2.png"), 6.8, "N2 diagram of the eight first-level functions from Innoslate (model 10.8).")
para(d, f"**Reading the diagram, source function by source function.** {T('n2flows')} summarises the same cells and adds the output that leaves F.2.", after=2)
n2rows = [
    ["F.2.0 Operate Service", "F.2.4: IO.21 Retention Directive (control). F.2.7: IO.22 Service Status Record. Outside F.2: IO.10 Release Configuration.", "Operations tells records when to delete everything at retirement and gives the pilot summary its availability data. It sends nothing into the UC.1 chain."],
    ["F.2.1 Collect Commute Entries", "F.2.2, F.2.3, F.2.4, F.2.5: IO.12 Validated Commute Entry. F.2.4: IO.13 Record Deletion Request.", "The widest fan-out: one validated entry feeds four functions, which is why intake is the first gate of the whole chain. IO.3 (the entry form) leaves to the Commuter."],
    ["F.2.2 Manage Reference Data", "F.2.3: IO.14 Emission Factor Record (control). F.2.5: IO.15 Trip Weather Forecast. Outside F.2: IO.4 Weather Request.", "Reference data governs the estimate through the factor and feeds the alternative rules through the forecast."],
    ["F.2.3 Estimate Commute Emissions", "F.2.4: IO.16 Trip Emission Estimate. F.2.6: IO.16 and IO.17 Estimation Method Disclosure (control).", "The estimate is stored by F.2.4 and shown by F.2.6; the method disclosure only goes to presentation, because only F.2.6 shows it."],
    ["F.2.4 Manage Commute Records", "F.2.6: IO.18 Commute Trend Summary. F.2.7: IO.19 Trip Record Extract.", "Records feed two consumers: the Commuter's own trend and the de-identified pilot extract."],
    ["F.2.5 Evaluate Feasible Alternatives", "F.2.6: IO.20 Alternative Recommendation.", "A single output into presentation."],
    ["F.2.6 Present Commute Results", "none (its only output, IO.3, leaves F.2 to the Commuter).", "A sink: it consumes results and produces the display."],
    ["F.2.7 Summarize Pilot Results", "none (its only output, IO.8, leaves F.2 to the Pilot Administrator).", "A sink for the pilot reporting chain."],
]
add_table("n2flows", "N2 diagram read row by row: internal flows leaving each function.", ["Source function (row)", "Flows to destination functions (cells)", "What it means"], n2rows, [1.6, 2.8, 2.7], size=8, first_bold=True)
para(d, "**Observations.** (1) All 13 filled cells lie above the diagonal. There is no feedback flow between first-level functions, so the functions can be built and tested in numeric order, and a change in a later function cannot alter an earlier function's inputs. (2) F.2.1 has the highest fan-out (four destination functions) and F.2.4 and F.2.6 the highest fan-in (three source functions each). Their requirements, 3.3.1.1, 3.3.1.4 and 3.3.1.6, are the ones to verify most carefully. (3) F.2.0 and F.2.7 reach the UC.1 chain only through F.2.4, so the UC.1 chain and the pilot reporting chain are loosely coupled; this supports changing either without the other. (4) F.2.5 depends on F.2.1 and F.2.2 only and does not depend on the estimate; the alternative rules use distance, weather and settings, as the OpsCon says. (5) The N2 reflects the trace edit that sends the validated entry to F.2.4 (cell F.2.1 to F.2.4 now holds IO.12 as well as IO.13).")

heading(d, "2.1.4 Architecture Evaluation", 3)
heading(d, "2.1.4.1 Tracing Analysis through OpsCon and IDEF", 4)
para(d, "The trace checks that the architecture supports the OpsCon. It follows the direct trace method of Section 6.2.1: the input/outputs of the UC.1 action diagram are numbered in order and found on the Metafunction (Part 1); the flows are then traced through the first-level IDEF0 with the hierarchical numbering n and n.m, using a sequence diagram that includes the internal interactions (Part 2); and the architecture is corrected where the trace breaks (Part 3). The whole trace uses one baseline diagram: the first-level IDEF0 of model 10.7 as exported on 24 October 2026.")
heading(d, "2.1.4.1.1 OpsCon (Operational Concept) and Metafunction Trace", 5)
para(d, "**OpsCon.** A student or employee opens EcoCommute to record a commute. The application presents fields for mode, distance, date, occupancy and feasibility settings. The commuter submits the trip. EcoCommute accepts the entry, retrieves a reviewed emission factor from its internal store, requests the weather forecast and receives it, estimates CO2e, stores the record, updates trends and compares the trip with an always-drive baseline. It evaluates transparent rules and displays the impact, trend, baseline difference and one feasible alternative. The commuter reviews the result and keeps the choice of future travel mode.")
para(d, f"**Action diagram.** The model's action diagram for UC.1 ({F('action')}) has 12 actions in three swimlanes (Commuter, EcoCommute, Open-Meteo) and six input/outputs. They are numbered 2 to 7 in the order in which they occur. Step 1 is the Reviewed Emission Factors (IO.6), a precondition that the action diagram does not show: factors are ingested before the scenario, so it is traced as step 1 and drawn dashed in the sequence diagram. In {F('action')} the places where the action diagram carries no input/output between two actions are the plain arrows from UC.1.3 to UC.1.4, UC.1.5 to UC.1.6 and UC.1.8 to UC.1.9 to UC.1.10 to UC.1.11.")
add_figure("action", os.path.join(DIA, "innoslate_action_diagram_numbered.png"), 7.0, "UC.1 action diagram from Innoslate, with its six input/outputs numbered 2 to 7 in order of occurrence.")
para(d, f"**Trace over the Metafunction.** {F('metatrace')} marks the numbered flows on the Innoslate Metafunction IDEF0. The six flows of the action diagram and the precondition use IO.1 to IO.6. IO.1 carries two numbers (2 and 4) and IO.3 carries two (3 and 7), because the Metafunction bundles related flows. {T('actmap')} maps each numbered interaction to its Metafunction input/output.", size=10)
add_figure("metatrace", os.path.join(DIA, "innoslate_metafunction_trace.png"), 6.4, "Trace of the UC.1 input/outputs over the Metafunction IDEF0 (Innoslate screenshot with the traced lines in red). Black lines are not used by UC.1.")
actmap = [
    ["1", "Reviewed Emission Factors", "X.4 to C.0", "precondition", "IO.6 Reviewed Emission Factors", "Found"],
    ["2", "Entry Request", "X.1 to C.1", "UC.1.1", "IO.1 Commuter Entries and Requests (more specific)", "Found"],
    ["3", "Entry Form", "C.1 to X.1", "UC.1.2", "IO.3 Commute Result and Advice", "**Not found** before the trace; found after edit 1"],
    ["4", "Commute Request", "X.1 to C.2", "UC.1.3, UC.1.4", "IO.1 Commuter Entries and Requests; settings on IO.2", "Found"],
    ["5", "Weather Query", "C.3 to X.3", "UC.1.6", "IO.4 Weather Request (same flow, different name)", "Found"],
    ["6", "Weather Forecast", "X.3 to C.3", "UC.1.7", "IO.5 Weather Forecast Data (same flow, different name)", "Found"],
    ["7", "Commute Impact Result", "C.1, C.4 to X.1", "UC.1.11, UC.1.12", "IO.3 Commute Result and Advice (more specific)", "Found"],
]
add_table("actmap", "Action diagram input/outputs numbered in sequence and found on the Metafunction.", ["No.", "Interaction", "From to", "UC.1 action", "Metafunction input/output", "Result"], actmap, [0.35, 1.4, 0.9, 0.9, 2.5, 1.05], size=8)
para(d, "Five of the six interactions, and the precondition, are accounted for on the Metafunction. The Entry Form is the exception. The Metafunction had no output that could carry it to the Commuter: IO.3 is defined as the result and advice and was generated only by F.2.6. This is the first inconsistency (Finding 1 in Section 2.1.4.1.3). The names Weather Query and Weather Forecast differ from IO.4 and IO.5 but name the same flows, so no change was needed.")

heading(d, "2.1.4.1.2 IDEF0 Functional Analysis: Trace through the First Level", 5)
para(d, f"**Sequence diagram.** The action diagram shows external interactions only, because the SoI is not decomposed in it. The model's own sequence diagram for UC.1 ({F('innoseq')}) cannot serve as the tracing diagram. It shows only six of the twelve actions (UC.1.1, 1.2, 1.4, 1.6, 1.7 and 1.11) and leaves out UC.1.3, 1.5, 1.8, 1.9, 1.10 and 1.12. Its lifelines are the components C.1 to C.4 and the external assets, not the first-level functions, so it cannot show the flows between F.2.1 and F.2.6. Its arrows carry the action names but not the input/output names, and they are not numbered. We therefore redrew the sequence over the first-level functions ({F('seq')}); it is the only figure in this report drawn outside Innoslate.")
add_figure("innoseq", os.path.join(DIA, "innoslate_sequence_diagram.png"), 5.6, "Sequence diagram of UC.1 from Innoslate (model 10.8): incomplete for tracing, see text.")
para(d, f"{F('seq')} adds the internal interactions between the first-level functions. It has 14 numbered interactions with unique names and numbers: external ones are numbered 1 to 7 (red) and internal ones n.m (blue), where n is the external step they follow. Message 4.1 reaches four functions and 6.2 reaches two, so they are drawn with one arrow head per receiver.")
add_figure("seq", os.path.join(DIA, "fig_sequence_portrait.png"), 7.0, "Sequence diagram of UC.1 over the first-level functions. Each interaction is numbered and named once; the IDEF0 flow is in brackets.")
para(d, f"**Trace over the first-level IDEF0.** {F('trace')} repeats the numbers on the first-level IDEF0. Red lines carry flows down from the Metafunction and keep the plain numbers 1 to 7; blue lines are internal flows with the numbers n.m; black lines are not used by UC.1. Every number sits on the line it names and in the same colour. {T('seqmap')} maps each numbered step of the sequence diagram to its IDEF0 flow and path.", size=10, after=2)
add_figure("trace", os.path.join(DIA, "innoslate_first_level_trace.png"), 7.0, "UC.1 traced through the first-level IDEF0 (Innoslate screenshot, model 10.8): carried-down flows in red, internal flows in blue, unused flows in black.")
seqmap = [
    ["1", "Reviewed Emission Factors", "IO.6", "Factor Publisher (F.5) to F.2.2 (control)"],
    ["2", "Entry Request", "IO.1", "Commuter (F.1) to F.2.1"],
    ["3", "Entry Form", "IO.3", "F.2.1 to Commuter (F.1)  [F.2.1 added as a source by edit 1]"],
    ["4", "Commute Request", "IO.1, IO.2", "Commuter (F.1) to F.2.1; settings (IO.2) also reach F.2.4 and F.2.5 as controls"],
    ["4.1", "Validated Commute Entry", "IO.12", "F.2.1 to F.2.2, F.2.3, F.2.4 [edit 2], F.2.5"],
    ["4.2", "Emission Factor Record", "IO.14", "F.2.2 to F.2.3 (control)"],
    ["5", "Weather Request", "IO.4", "F.2.2 to Open-Meteo (F.4)"],
    ["6", "Weather Forecast Data", "IO.5", "Open-Meteo (F.4) to F.2.2"],
    ["6.1", "Trip Weather Forecast", "IO.15", "F.2.2 to F.2.5"],
    ["6.2", "Trip Emission Estimate", "IO.16", "F.2.3 to F.2.4 and F.2.6"],
    ["6.3", "Estimation Method Disclosure", "IO.17", "F.2.3 to F.2.6 (control)"],
    ["6.4", "Commute Trend Summary", "IO.18", "F.2.4 to F.2.6"],
    ["6.5", "Alternative Recommendation", "IO.20", "F.2.5 to F.2.6"],
    ["7", "Commute Impact Result", "IO.3", "F.2.6 to Commuter (F.1)"],
]
add_table("seqmap", "Numbered steps of the sequence diagram mapped to the IDEF0 flows. Every numbered step appears once on the traced IDEF0.", ["Step", "Interaction", "IDEF0 flow", "Path on the first-level IDEF0"], seqmap, [0.45, 2.0, 0.85, 3.8], size=8.2)
para(d, "**Pencil test.** Following the pencil through the diagram: the Commuter's request (2) enters F.2.1, which answers with the form (3). The submitted request (4) enters F.2.1 again, and F.2.1 releases the validated entry (4.1). F.2.2 holds the factor (1) and the entry, and sends the factor record to F.2.3 (4.2) and the weather request to Open-Meteo (5); the forecast (6) returns to F.2.2, which passes it to F.2.5 (6.1). F.2.3, with entry and factor record, sends the estimate to F.2.4 and F.2.6 (6.2) and the disclosure to F.2.6 (6.3). F.2.4 stores the trip and sends the trend to F.2.6 (6.4); F.2.5 sends the alternative to F.2.6 (6.5). F.2.6 returns the result to the Commuter (7). The path branches after 4.1 and 6, and the branches run concurrently and rejoin at F.2.6. Each branch ends at F.2.6 without a jump, which is the exception the lab allows for concurrent branches. There is no function whose output appears without an input, and every function on the path (F.2.1 to F.2.6) receives what it needs from the previous step. The trace therefore passes the pencil test.")

heading(d, "2.1.4.1.3 Summarize the Work", 5)
para(d, f"{T('findings')} lists the errors and inconsistencies the process exposed and how each was resolved. Findings 1 and 2 changed the model; the others were resolved by documentation or are carried into the next chapter.")
findings = [
    ["1", "Entry Form (UC.1.2, step 3) had no output on the Metafunction or first level that could carry it to the Commuter. F.2.1 is described as presenting the form, but only F.2.6 generated IO.3.", "IDEF0 incomplete", "**Edit 1.** F.2.1 added as a second source of IO.3; IO.3's description now covers the entry form and validation feedback. IO.3 is a bundle, as the lab recommends for broad top-level flows."],
    ["2", "F.2.4 stores trips and builds weekly and monthly trends, but its only inputs were the deletion request and the estimate. The trip's mode, distance and date, held in the validated entry, did not reach it.", "IDEF0 gap", "**Edit 2.** IO.12 Validated Commute Entry is now also received by F.2.4 (cell F.2.1 to F.2.4 of the N2 diagram). Innoslate attaches the added arrow to the top of F.2.4 and does not let it be moved, so on the first-level IDEF0 and in this report IO.12 is a control of F.2.4: the validated entry fixes which trip record is stored."],
    ["3", "The action diagram has no input/output between UC.1.3 and 1.4 (same performer), UC.1.5 and 1.6, and UC.1.8, 1.9, 1.10 and 1.11. The internal interactions are not in it.", "Action diagram incomplete at first level", "No change to the action diagram: its performers are the components C.1 to C.4, not the first-level functions. The sequence diagram carries the internal interactions 4.1 to 6.5."],
    ["4", "Weather Query and Weather Forecast on the action diagram differ from Weather Request (IO.4) and Weather Forecast Data (IO.5); Entry Request, Commute Request and Commute Impact Result are more specific than IO.1 and IO.3.", "Naming", "No change. Names may differ with the level of detail. IO.1 carries steps 2 and 4 and IO.3 steps 3 and 7."],
    ["5", "Factor ingestion and recorded consent are preconditions of UC.1 and are not on the action diagram.", "Missing precondition", "Traced as step 1 (IO.6, dashed). Consent is assumed and not traced."],
    ["6", "UC.1.5 (performed by C.2) spans F.2.1 and F.2.2, and UC.1.11 (C.1 with C.4) is realized by F.2.6 alone.", "Allocation to form", "No change now. Resolved when first-level functions are allocated to mechanisms (risk TR.9)."],
]
add_table("findings", "Errors, inconsistencies and observations exposed by the trace.", ["No.", "Finding", "Type", "Resolution"], findings, [0.35, 2.9, 1.0, 2.85], size=8)
para(d, f"**What changed in the model.** {T('changes')} lists the model changes: two relationship additions and one description. They are model version 10.8 and can be applied by hand in Innoslate in a few minutes (Appendix B).")
changes = [
    ["Edit 1", "IO.3 Commute Result and Advice", "Add generated by F.2.1 (F.2.6 stays); extend description", "First-level IDEF0, N2 not affected, sequence step 3"],
    ["Edit 2", "IO.12 Validated Commute Entry", "Add received by F.2.4", "First-level IDEF0, N2 cell F.2.1 to F.2.4, sequence step 4.1"],
]
add_table("changes", "Model changes made as a result of the trace (10.7 to 10.8).", ["Edit", "Entity", "Change", "Where it shows"], changes, [0.6, 2.1, 2.4, 2.0], size=8.2)
para(d, "**Insights.** (1) UC.1 engages six of the eight first-level functions (F.2.1 to F.2.6) and 13 of the 22 input/outputs: IO.1 to IO.6 from the Metafunction and seven internal flows. F.2.0 and F.2.7 and nine input/outputs (IO.7 to IO.11, IO.13, IO.19, IO.21 and IO.22) are idle; they belong to UC.2 and to the operations actions, so the next trace should cover UC.2. (2) The two errors were both found at the boundary between views, not within one view: the form was in the action diagram but not the IDEF0, and the trip data was on the IDEF0 only as an implied part of the estimate. Tracing is the only check that crosses these views. (3) The internal flows needed to complete the trace are exactly the first-level flows already in the IDEF0, apart from the two edits; the first-level decomposition was therefore sound. (4) Bundling input/outputs (IO.1, IO.3) kept the Metafunction small without hiding any step: each bundle carries two numbers. (5) The trace has no feedback loops, consistent with the N2 diagram.")

# ================================================================== 3 System Requirements Management
heading(d, "3 System Requirements Management", 1)
heading(d, "3.1 Requirements Traceability", 2)
heading(d, "3.1.1 System Functional Requirements", 3)
para(d, f"{T('freq')} lists the system functional requirements: the single top-level requirement 3.3.1 and its eight first-level children 3.3.1.1 to 3.3.1.8, one per first-level function. Each is a single \"shall\" statement. Each maps to a higher-level requirement: the parent 3.3.1 (and through it group 3.3 and the root 3 System Requirements) and the stakeholder requirements (SRD 2.x.x) that it traces from. The rationale, the measure of performance (MOP) and the verification criteria are stored in the model with each requirement; the table reproduces them. The requirement statements follow the form produced by Innoslate's Generate Requirements script, so the statements stay general and the measurable target is carried by the MOP.")
fr = []
fun_of = {"3.3.1": "F.2 Provide Commute Impact Decision Support", "3.3.1.1": "F.2.1 Collect Commute Entries", "3.3.1.2": "F.2.2 Manage Reference Data",
          "3.3.1.3": "F.2.3 Estimate Commute Emissions", "3.3.1.4": "F.2.4 Manage Commute Records", "3.3.1.5": "F.2.5 Evaluate Feasible Alternatives",
          "3.3.1.6": "F.2.6 Present Commute Results", "3.3.1.7": "F.2.7 Summarize Pilot Results", "3.3.1.8": "F.2.0 Operate Service"}
for num, fn in fun_of.items():
    r = parse_req(num)
    parent = "3.3 Functional Requirements (root 3)" if num == "3.3.1" else "3.3.1 Provide commute impact decision support"
    fr.append([num, r["stmt"], fn, parent + "\nSRD " + r["srd"], r["rationale"], r["mop"], r["ver"]])
add_table("freq", "System functional requirements with satisfying function, higher-level requirement, rationale, MOP and verification criteria.",
          ["ID", "Requirement (description)", "Satisfied by", "Higher-level requirement", "Rationale", "Measure of performance (MOP)", "Verification criteria"], fr,
          [0.5, 1.0, 0.85, 1.05, 1.35, 1.2, 1.15], size=6.8, first_bold=True)
heading(d, "3.1.2 Function-to-Requirement Traceability", 3)
para(d, f"{F('matrix')} is the traceability matrix. Its rows are only functions (F.0 to F.6 and F.2.0 to F.2.7); its columns are the system requirements and their groups: input/output (3.1.x), external interface (3.2.x), functional (3.3.x) and system-wide (3.4.x). An X is a satisfied by relationship in the model. The model does not yet relate the input/output requirements to first-level functions, so {T('cover')} adds, with the prefix o:, the input/output requirements that each first-level function sends or receives according to the first-level IDEF0. {T('cover')} then explains the coverage of each of the 15 functions.")
add_figure("matrix", os.path.join(DIA, "innoslate_traceability_matrix.png"), 7.1, "Requirements traceability matrix from Innoslate: functions (rows) against system requirements (columns). X = satisfied by.")
cover = [
    ["F.0 Metafunction", "none", "Structural parent of all functions.", "No requirement is allocated to F.0 by design; requirements attach to F.2 and its children, and to the external functions through interfaces."],
    ["F.1 Use EcoCommute Decision Support", "3.2.1, 3.2.2", "Commuter side of the web entry form and result view.", "Interface requirements only; the Commuter's behavior is outside our design scope."],
    ["F.2 Provide Commute Impact Decision Support", "3.1.1.1 to 3.1.1.7, 3.1.2.1 to 3.1.2.4, 3.2.1 to 3.2.8, 3.3.1, 3.4.1.1 to 3.4.6.1 (31)", "Accept and provide every top-level input/output, support each interface, deliver the top-level function, and stay within the system-wide thresholds.", "Carries all input/output, interface and system-wide requirements; the eight first-level functions share the functional ones."],
    ["F.2.0 Operate Service", "3.3.1.8; o: 3.1.1.6, 3.1.1.7, 3.1.2.4", "Deploy, monitor and retire the service within 99.5% availability and with no identifiable records left.", "Monitoring and retirement rehearsal verify it; it receives rules and hosting status and sends the release configuration."],
    ["F.2.1 Collect Commute Entries", "3.3.1.1; o: 3.1.1.1, 3.1.1.2, 3.1.2.1", "Present the form and feedback, validate entries and enforce consent so that a median entry takes 30 s or less.", "Timed usability session verifies it; it receives the entries and settings and sends the form."],
    ["F.2.2 Manage Reference Data", "3.3.1.2; o: 3.1.1.3, 3.1.1.4, 3.1.2.2", "Use only reviewed factors with provenance and obtain the forecast.", "Inspection of factor records verifies it; it handles the factors, forecast data and weather request."],
    ["F.2.3 Estimate Commute Emissions", "3.3.1.3", "Reproduce reference-case CO2e exactly (target 0 g difference after rounding).", "Reference-case test run verifies it; it has no top-level input/output of its own."],
    ["F.2.4 Manage Commute Records", "3.3.1.4; o: 3.1.1.2", "Store trips under consent, expose them to no other role and delete on request.", "Role access matrix and deletion test verify it; it uses the settings as a control."],
    ["F.2.5 Evaluate Feasible Alternatives", "3.3.1.5; o: 3.1.1.2", "Never recommend a mode the Commuter marked infeasible.", "Rule test over all excluded-mode combinations verifies it; it uses the settings as a control."],
    ["F.2.6 Present Commute Results", "3.3.1.6; o: 3.1.2.1", "Show the result within 3 s (TBD) with no unsupported facts.", "Timing on three devices and inspection of 50 explanations verify it; it delivers the result and advice."],
    ["F.2.7 Summarize Pilot Results", "3.3.1.7; o: 3.1.1.5, 3.1.1.6, 3.1.2.3", "Report no group under 5 (TBD) and no identifiers.", "Seeded aggregation test verifies it; it receives the summary request and rules and issues the approved summary."],
    ["F.3 Administer Pilot Program", "3.2.6, 3.2.7", "Authenticated dashboard and approved aggregate export.", "Interface requirements only."],
    ["F.4 Provide Weather Forecasts", "3.2.3, 3.2.4", "HTTPS API request to and structured response from Open-Meteo.", "Interface requirements only."],
    ["F.5 Publish Emission Factors", "3.2.5", "Reviewed file or source-document ingestion.", "Interface requirement only."],
    ["F.6 Provide Hosting Service", "3.2.8", "Provider deployment and status interface.", "Interface requirement only."],
]
add_table("cover", "Requirements covered by each function, what the function must do, and a summary of coverage.", ["Function", "Requirements covered", "What the function must do", "Summary"], cover, [1.4, 1.5, 2.1, 2.1], size=7.8, first_bold=True)
para(d, "**Observations.** (1) Every one of the 39 requirements has at least one satisfying function; no requirement is orphaned and no function except F.0 is without a requirement. (2) Each first-level function has exactly one functional requirement, so the eight children mirror the eight functions one to one. (3) The input/output requirements (3.1.x) are satisfied only at F.2 in the model. The o: entries in " + T('cover') + " show which first-level function carries each one, and recording them as satisfied by relationships is a model task for the next chapter. (4) System-wide requirements (3.4.x) are satisfied at F.2 as a whole; "+T('nfr')+" names the functions they influence most. (5) The external functions carry only interface requirements, as the Metafunction implies.")

heading(d, "4 Summary", 1)
para(d, f"We built and traced the first-level function architecture of EcoCommute. The architecture has 15 functions on three levels (Metafunction, six top-level functions and eight first-level functions of F.2) and 22 input/outputs, none of which is a tunnel ({F('hier')}, {F('first')}). The N2 diagram shows 13 function-to-function connections, all above the diagonal, so the first-level functions have no feedback loops ({F('n2')}).")
para(d, f"We traced the nominal scenario UC.1 in two stages. Its six action-diagram input/outputs were numbered and found on the Metafunction, and 14 numbered interactions were traced through the first-level IDEF0 in the pencil test ({F('metatrace')}, {F('seq')}, {F('trace')}). The trace exposed two inconsistencies: the entry form had no output to carry it, and the records function did not receive the trip details. Both were fixed with two relationship edits (model 10.8). The remaining findings were resolved by documentation or deferred to allocation. The functional requirements (3.3.1 and 3.3.1.1 to 3.3.1.8) are stated with rationale, MOP and verification criteria, and all 39 system requirements trace to at least one function ({F('matrix')}).")
para(d, "**Next steps.** (1) Allocate the first-level functions to mechanisms and resolve findings 6 and TR.9. (2) Trace UC.2 and the operations actions through F.2.0 and F.2.7. (3) Record the o-marked allocations of input/output requirements as satisfied by relationships. (4) Confirm the 3-second and group-size-5 values with the sponsor.")

# ================================================================== Appendices
heading(d, "Appendix A. Model Assessment Check", 1)
para(d, f"{T('tracecrit')} checks the work against the Tracing Assessment Criteria and {T('idefcrit')} against the IDEF Model Assessment Criteria.")
tc = [
    ["Action diagram input/outputs are accounted for on the Metafunction", f"{T('actmap')}: all six found; Entry Form after edit 1.", "Met"],
    ["Every interaction on the sequence diagram is numbered in order and unique", f"{F('seq')}: 14 interactions, numbers 1 to 7 and n.m, no duplicate name.", "Met"],
    ["Sequence diagram has the complete set of interactions for a continuous flow (pencil test)", "Internal interactions 4.1 to 6.5 added; pencil test narrative in 2.1.4.1.2.", "Met"],
    ["All team members used the same legible IDEF0", f"One baseline (export of 24 Oct 2026) and one traced figure, {F('trace')}.", "Met"],
    ["No breaks in flow through the first level", "Branches end at F.2.6 without a jump.", "Met"],
    ["IDEF0 flows recolored, numbered and placed without ambiguity", f"{F('metatrace')} and {F('trace')}: red and blue lines, badge on the line it names.", "Met"],
    ["Every numbered IDEF0 step maps to the sequence diagram", f"{T('seqmap')}: 14 of 14.", "Met"],
    ["Figure included with explanation of errors and insights", f"Section 2.1.4.1.3, {T('findings')}.", "Met"],
]
add_table("tracecrit", "Tracing Assessment Criteria check.", ["Criterion", "Evidence", "Status"], tc, [3.0, 3.1, 1.0], size=8.2)
ic = [
    ["All functions have verb-oriented names", f"{T('funcs')} and {T('purpose')}; names such as Collect Commute Entries.", "Met"],
    ["Functions are numbered with a clear convention", "F.0, F.1 to F.6, F.2.0 to F.2.7.", "Met"],
    ["All input/outputs have noun-oriented names", T('ios') + ".", "Met"],
    ["Every input/output has a source and a destination (no tunnels)", f"{T('ios')}; {F('first')} and {F('metatrace')}.", "Met"],
    ["Each function has appropriate input/outputs", T('icom1') + ".", "Met"],
    ["Internal interactions give continuity for the covered use cases", "Trace in Section 2.1.4.1.2.", "Met for UC.1; UC.2 pending"],
    ["Each Metafunction box has exactly one mechanism from the Universe hierarchy", T('metaicom') + ".", "Met"],
    ["No mechanisms on the first-level IDEF0", "First-level functions are not allocated.", "Met (as expected)"],
    ["Diagram content is legible", f"{F('first')}, {F('n2')}, {F('metatrace')} and {F('trace')} are full-width figures; the cell text of the N2 diagram is also written out in {T('n2flows')}.", "Met"],
]
add_table("idefcrit", "IDEF Model Assessment Criteria check.", ["Criterion", "Evidence", "Status"], ic, [3.0, 3.1, 1.0], size=8.2)

heading(d, "Appendix B. Applying the Model Edits in Innoslate", 1)
para(d, "The patched export docs/model/EcoCommute_10.8_C6-C7_trace_update.xml contains both edits and can be imported. The same two edits can be made by hand in model 10.7.")
bullet(d, "**Edit 1.** Open IO.3 Commute Result and Advice. In its relationships add generated by: F.2.1 Collect Commute Entries, and keep F.2.6. Replace the description with: \"Entry form and validation feedback presented to the Commuter before the result (from F.2.1), then estimated CO2e, trend, baseline difference, method disclosure, one feasible alternative and plain-language explanation (from F.2.6).\"")
bullet(d, "**Edit 2.** Open IO.12 Validated Commute Entry. In its received by relationship add F.2.4 Manage Commute Records.")
bullet(d, "Open the first-level IDEF0 of F.2 and move the new arrow labels if the boxes resize.")
bullet(d, "**Restored links.** The 24 October export had lost 42 relationships that model 10.7 in the repository still had: the 20 related to links (and their inverses) between risks TR.1 to TR.9 and the requirements they threaten, and the traced to link from F.0 Metafunction to the root requirement 3 System Requirements. They are restored in the 10.8 file; if you edit by hand, re-add them from the repository copy of 10.7.")

header_footer(d, "EcoCommute Team 6  |  Architecture Definition: Function (Chapters 6 and 7)")
d.save(OUT)
assert _fig_seen == FIG_ORDER and _tab_seen == TAB_ORDER, (_fig_seen, _tab_seen)
print("saved", OUT)
